"""Conectores de datos: CSV, Excel, base SQL y API. Devuelven listas de diccionarios (una fila por registro).

Solo se lee: nunca se escribe en el sistema de inventario de la empresa. Para SQL, ademas de exigir una consulta SELECT,
la conexion se abre en modo de solo lectura donde el motor lo permite. Las credenciales (URL de la base y token de la API)
se leen de variables de entorno y no se guardan en la configuracion.
"""
import csv
import json
import os
import pathlib
import re
import sqlite3
import urllib.parse
import urllib.request

from . import red


class ErrorFuente(Exception):
    """La fuente de datos no se pudo leer; el mensaje explica que revisar."""


def leer_csv(ruta, separador=","):
    if not ruta or not os.path.isfile(ruta):
        raise ErrorFuente("No se encontro el archivo CSV: %s" % ruta)
    with open(ruta, encoding="utf-8-sig", newline="") as f:
        return [dict(fila) for fila in csv.DictReader(f, delimiter=separador or ",") if any((v or "").strip() for v in fila.values() if isinstance(v, str))]


def leer_excel(ruta, hoja):
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise ErrorFuente("Para leer Excel instala openpyxl (pip install openpyxl)") from None
    if not ruta or not os.path.isfile(ruta):
        raise ErrorFuente("No se encontro el libro de Excel: %s" % ruta)
    libro = load_workbook(ruta, read_only=True, data_only=True)
    try:
        if hoja not in libro.sheetnames:
            raise ErrorFuente("El libro no tiene la hoja '%s'. Hojas: %s" % (hoja, ", ".join(libro.sheetnames)))
        filas = libro[hoja].iter_rows(values_only=True)
        enc = next(filas, None)
        if not enc:
            return []
        nombres = [str(h).strip() if h is not None else "" for h in enc]
        out = []
        for fila in filas:
            if fila is None or all(v is None or (isinstance(v, str) and not v.strip()) for v in fila):
                continue
            out.append({n: v for n, v in zip(nombres, fila) if n})
        return out
    finally:
        libro.close()


_SOLO_SELECT = re.compile(r"\s*(?:--[^\n]*\n\s*)*(select|with)\b", re.IGNORECASE)


def validar_consulta(consulta):
    """Defensa en profundidad: solo una consulta SELECT (o WITH ... SELECT), sin varias sentencias."""
    q = (consulta or "").strip().rstrip(";").strip()
    if not q:
        raise ErrorFuente("La consulta SQL esta vacia")
    if not _SOLO_SELECT.match(q):
        raise ErrorFuente("Solo se permiten consultas SELECT: el agente no escribe en la base de datos")
    if ";" in q:
        raise ErrorFuente("La consulta SQL debe ser una sola sentencia")
    return q


def leer_sql(url, consulta):
    """url: sqlite:///ruta, mysql://usuario:clave@host:puerto/base o postgresql://usuario:clave@host:puerto/base."""
    q = validar_consulta(consulta)
    if not url:
        raise ErrorFuente("Falta la URL de la base de datos (variable de entorno)")
    partes = urllib.parse.urlsplit(url)
    esquema = partes.scheme.lower()
    if esquema == "sqlite":
        ruta = urllib.parse.unquote(url[len("sqlite:///"):] if url.lower().startswith("sqlite:///") else partes.path)
        if not os.path.isfile(ruta):
            raise ErrorFuente("No se encontro la base SQLite: %s" % ruta)
        con = sqlite3.connect(pathlib.Path(ruta).resolve().as_uri() + "?mode=ro", uri=True)
        cierre = con.close
        cur = con.cursor()
    elif esquema in ("mysql", "mariadb"):
        try:
            import pymysql
        except ImportError:
            raise ErrorFuente("Para MySQL instala pymysql (pip install pymysql)") from None
        con = pymysql.connect(host=partes.hostname, port=partes.port or 3306, user=urllib.parse.unquote(partes.username or ""), password=urllib.parse.unquote(partes.password or ""),
                              database=partes.path.lstrip("/"), connect_timeout=15)
        cierre = con.close
        cur = con.cursor()
        cur.execute("START TRANSACTION READ ONLY")
    elif esquema in ("postgresql", "postgres"):
        try:
            import psycopg2 as pg
            con = pg.connect(url, connect_timeout=15)
            con.set_session(readonly=True)
        except ImportError:
            try:
                import psycopg as pg3
                con = pg3.connect(url, connect_timeout=15)
                con.read_only = True
            except ImportError:
                raise ErrorFuente("Para PostgreSQL instala psycopg2 (pip install psycopg2-binary) o psycopg") from None
        cierre = con.close
        cur = con.cursor()
    else:
        raise ErrorFuente("Base de datos no soportada: '%s'. Usa sqlite, mysql o postgresql" % esquema)
    try:
        cur.execute(q)
        columnas = [d[0] for d in cur.description or []]
        return [dict(zip(columnas, fila)) for fila in cur.fetchall()]
    finally:
        cierre()


def leer_api(url, ruta_datos="", token=None, timeout=60):
    """GET que devuelve JSON: una lista de registros, o un objeto que trae la lista en el campo `ruta_datos`."""
    if not url or not re.match(r"https?://", url, re.IGNORECASE):
        raise ErrorFuente("La URL de la API debe empezar con http:// o https://")
    encabezados = {"Accept": "application/json"}
    if token:
        encabezados["Authorization"] = "Bearer %s" % token
    try:
        red.exigir_https_si_hay_clave(url, bool(token))
        with red.abrir(urllib.request.Request(url, headers=encabezados, method="GET"), timeout) as r:
            carga = json.loads(r.read().decode("utf-8"))
    except red.ErrorRed as e:
        raise ErrorFuente(str(e)) from None
    except Exception as e:  # noqa: BLE001 - se explica el motivo sin exponer encabezados ni token
        raise ErrorFuente("No se pudo leer la API (%s)" % type(e).__name__) from None
    if isinstance(carga, dict) and ruta_datos:
        carga = carga.get(ruta_datos)
    if isinstance(carga, dict):
        carga = [carga]
    if not isinstance(carga, list):
        raise ErrorFuente("La respuesta de la API no es una lista de registros" + (" (campo '%s')" % ruta_datos if ruta_datos else ""))
    return [f for f in carga if isinstance(f, dict) and not f.get("error")]


def leer_fuente(cfg):
    """Lee movimientos y productos segun cfg['fuente'] (csv, excel, sql o api). Devuelve (filas_mov, filas_mae)."""
    fuente = (cfg.get("fuente") or "csv").lower()
    if fuente == "csv":
        sep = cfg.get("separador") or ","
        return leer_csv(cfg.get("archivo_movimientos"), sep), leer_csv(cfg.get("archivo_productos"), sep) if cfg.get("archivo_productos") else []
    if fuente == "excel":
        return leer_excel(cfg.get("archivo_excel"), cfg.get("hoja_movimientos") or "Movimientos"), leer_excel(cfg.get("archivo_excel"), cfg.get("hoja_productos") or "Productos")
    if fuente == "sql":
        url = os.environ.get(cfg.get("sql_url_env") or "INVENTARIOS_SQL_URL", "")
        return leer_sql(url, cfg.get("sql_movimientos")), (leer_sql(url, cfg.get("sql_productos")) if cfg.get("sql_productos") else [])
    if fuente == "api":
        token = os.environ.get(cfg.get("api_token_env") or "INVENTARIOS_API_TOKEN", "") or None
        ruta = cfg.get("api_ruta_datos") or ""
        return leer_api(cfg.get("api_url_movimientos"), ruta, token), (leer_api(cfg.get("api_url_productos"), ruta, token) if cfg.get("api_url_productos") else [])
    raise ErrorFuente("Fuente desconocida: '%s'. Usa csv, excel, sql o api" % fuente)
